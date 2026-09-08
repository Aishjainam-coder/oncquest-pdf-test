import json
import docx

json_path = r'extracted_jsons/TestReport_SAROJ DEVI (OQG2604250052)_2600133939_1eb3f888-f97e-4181-b84d-d9536de7af26.json'
docx_path = r'output/TestReport_SAROJ DEVI (OQG2604250052)_2600133939_1eb3f888-f97e-4181-b84d-d9536de7af26_report.docx'

with open(json_path, 'r', encoding='utf-8') as f:
    data = json.load(f)

doc = docx.Document(docx_path)
docx_full = '\n'.join([p.text for p in doc.paragraphs] + [c.text for t in doc.tables for r in t.rows for c in r.cells])

doc_pages = data.get('document', {}).get('pages', [])

missing_details = []
for p_idx, page in enumerate(doc_pages):
    p_num = page.get('page_number', p_idx + 1)
    elements = page.get('elements', [])
    for el_idx, el in enumerate(elements):
        el_type = el.get('type')
        text = el.get('text', '') or el.get('title', '')
        # Check if table
        if el_type == 'table':
            headers = el.get('headers', [])
            rows = el.get('rows', [])
            sample = ' '.join(str(h) for h in headers[:3])
            # check if sample in docx_full
            if sample and sample[:25] not in docx_full:
                missing_details.append((p_num, el_type, sample, el))
        elif el_type in ('paragraph', 'heading', 'text_block', 'banner'):
            if text and len(text.strip()) > 5:
                # check if first 25 chars in docx
                s = text.strip()[:25]
                if s not in docx_full:
                    missing_details.append((p_num, el_type, text.strip()[:60], el))
        elif el_type in ('box', 'content_section'):
            content = el.get('content', '') or el.get('text', '') or el.get('title', '')
            if content and len(content.strip()) > 5:
                s = content.strip()[:25]
                if s not in docx_full:
                    missing_details.append((p_num, el_type, content.strip()[:60], el))

print(f"Total elements in document.pages: {sum(len(p.get('elements', [])) for p in doc_pages)}")
print(f"Elements detected as missing/not-matching: {len(missing_details)}")
print("\n--- Details of missing elements ---")
for p_num, el_type, snippet, el in missing_details:
    print(f"Page {p_num} [{el_type}] bbox={el.get('bbox')}: {repr(snippet)}")
