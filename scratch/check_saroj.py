import json
import docx
import sys

json_path = r'extracted_jsons/TestReport_SAROJ DEVI (OQG2604250052)_2600133939_1eb3f888-f97e-4181-b84d-d9536de7af26.json'
docx_path = r'output/TestReport_SAROJ DEVI (OQG2604250052)_2600133939_1eb3f888-f97e-4181-b84d-d9536de7af26_report.docx'

with open(json_path, 'r', encoding='utf-8') as f:
    data = json.load(f)

doc = docx.Document(docx_path)

docx_text = set()
docx_raw_text = []
for p in doc.paragraphs:
    t = p.text.strip()
    if t:
        docx_raw_text.append(t)
        for w in t.split():
            docx_text.add(w.lower())
for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            t = cell.text.strip()
            if t:
                docx_raw_text.append(t)
                for w in t.split():
                    docx_text.add(w.lower())

all_docx_blob = " ".join(docx_raw_text).lower()

print(f"Total pages in JSON: {len(data.get('pages', []))}")
print(f"Total sections in JSON: {len(data.get('sections', []))}")
print(f"Total tables in JSON: {len(data.get('tables', []))}")

# Check which items/blocks from JSON are missing in DOCX
missing_blocks = []
for p_idx, page in enumerate(data.get('pages', [])):
    p_num = page.get('page_number', p_idx + 1)
    for tb in page.get('text_blocks', []):
        txt = tb.get('text', '').strip()
        if len(txt) > 20:
            sample = txt[:40].lower()
            if sample not in all_docx_blob:
                missing_blocks.append((p_num, 'text_block', txt))
    for tbl in page.get('tables', []):
        headers = tbl.get('headers', [])
        rows = tbl.get('rows', [])
        tbl_text = " ".join([str(x) for x in headers] + [str(c) for r in rows for c in r])
        sample = tbl_text[:40].lower()
        if sample and sample not in all_docx_blob:
            missing_blocks.append((p_num, 'table', tbl_text[:80]))
    for cs in page.get('content_sections', []):
        cs_title = cs.get('title', '')
        cs_text = cs.get('content', '') or cs.get('text', '')
        if cs_text and len(cs_text) > 20 and cs_text[:40].lower() not in all_docx_blob:
            missing_blocks.append((p_num, 'content_section', f"{cs_title}: {cs_text[:60]}"))

print(f"Missing items count: {len(missing_blocks)}")
for p_num, typ, snippet in missing_blocks[:20]:
    print(f"[Page {p_num}] [{typ}] -> {repr(snippet)}")
