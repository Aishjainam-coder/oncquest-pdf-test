import re
import copy
import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def is_end_of_report_text(text: str) -> bool:
    if not isinstance(text, str):
        return False
    return bool(re.search(r'end\s+of\s+report', text, re.IGNORECASE))

def is_test_name_text(text: str) -> bool:
    if not isinstance(text, str):
        return False
    cleaned = text.strip()
    if len(cleaned) < 4 or len(cleaned) > 150:
        return False
        
    exclude_prefixes = (
        "clinical indication", "sample description", "report highlights", 
        "key findings", "test results", "tier ", "case id", "sample type", 
        "name :", "date & time", "bill. loc", "ref. by", "report version",
        "qr code", "page ", "salient features", "clinical suspicion",
        "dr.", "laboratory", "oncquest", "result summary", "methodology",
        "test description", "extraction", "the performance", "analyze",
        "test name:"
    )
    cleaned_lower = cleaned.lower()
    for ex in exclude_prefixes:
        if cleaned_lower.startswith(ex):
            return False
            
    patterns = [
        r'^(?:Breast\s+and\s+Ovarian\s+(?:Cancer\s+)?Extended\s+Panel\s*[-–]\s*Liquid\s+Biopsy\s+Assay)$',
        r'^(?:(?:Liquidseq|Solidseq|Brainseq)[\w\s\(\)/&,\.\+\-]*?(?:Genomic\s+Profiling\s+Panel|Panel|Assay)(?:\s*[-–]\s*Advance)?(?:\s*\([^)]*\))?(?:\s+on\s+(?:the\s+)?Illumina\s+[\w\s-]+\s+Platform)?)$',
        r'^(?:Whole\s+Exome\s+Sequencing(?:\s+on\s+(?:the\s+)?Illumina\s+[\w\s-]+\s+Platform)?)$',
        r'^(?:(?:Breast|Ovarian|Lung|Colorectal|Pancreatic|Hereditary|Comprehensive|Focus|Actionable|Fusion|Somatic|Germline)[\w\s/&,–\-\(\)\.\+]+?(?:Genomic\s+Profiling\s+Panel|Extended\s+Panel|Profiling\s+Panel|Biopsy\s+Assay|Exome\s+Sequencing|Sequencing\s+Panel|Cancer\s+Panel|Gene\s+Panel|Profiling\s+Assay|Sequencing\s+Assay|Biopsy\s+Panel|NGS\s+Panel|Comprehensive\s+Panel)(?:\s*[-–]\s*Advance)?(?:\s*\([^)]*\))?(?:\s+on\s+(?:the\s+)?Illumina\s+[\w\s-]+\s+Platform)?)$',
        r'^(?:[A-Z\s]{4,}\s+PANEL(?:\s*[-–]\s*ADVANCE)?(?:\s*\([^)]*\))?)$',
    ]
    for pat in patterns:
        if re.match(pat, cleaned, re.IGNORECASE):
            return True
            
    return False

KNOWN_TEST_NAMES_REGEX = (
    r'(?:Breast\s+and\s+Ovarian\s+(?:Cancer\s+)?Extended\s+Panel\s*[-–]\s*Liquid\s+Biopsy\s+Assay|'
    r'(?:Liquidseq|Solidseq|Brainseq)[\w\s\(\)/&,\.\+\-]*?(?:Genomic\s+Profiling\s+Panel|Comprehensive\s+Panel|Cancer\s+Panel|Panel|Assay)(?:\s*[-–]\s*Advance)?(?:\s*\([^)]*\))?(?:\s+on\s+(?:the\s+)?Illumina\s+[\w\s-]+\s+Platform)?|'
    r'Whole\s+Exome\s+Sequencing(?:\s+on\s+(?:the\s+)?Illumina\s+[\w\s-]+\s+Platform)?|'
    r'(?:Breast|Ovarian|Lung|Colorectal|Pancreatic|Hereditary|Comprehensive|Actionable|Fusion|Somatic|Germline)[\w\s/&,–\-\(\)\.\+]*?(?:Genomic\s+Profiling\s+Panel|Extended\s+Panel|Profiling\s+Panel|Biopsy\s+Assay|Exome\s+Sequencing|Sequencing\s+Panel|Cancer\s+Panel|Gene\s+Panel|Profiling\s+Assay|Sequencing\s+Assay|Biopsy\s+Panel|NGS\s+Panel|Comprehensive\s+Panel)(?:\s*[-–]\s*Advance)?(?:\s*\([^)]*\))?(?:\s+on\s+(?:the\s+)?Illumina\s+[\w\s-]+\s+Platform)?)'
)

test_name_pattern = re.compile(
    r'(?:\bTEST\s+NAME\b(?!\s*[:=])|' + KNOWN_TEST_NAMES_REGEX + r')',
    re.IGNORECASE
)

highlight_span = '<span style="background-color: #ffff00; color: #000000; font-weight: bold; padding: 1px 4px; border-radius: 2px;">TEST NAME</span>'

def replace_test_name_in_html(html: str) -> str:
    if not isinstance(html, str):
        return html
    m = re.search(r'end\s+of\s+report', html, re.IGNORECASE)
    pre = html[:m.start()] if m else html
    post = html[m.start():] if m else ''

    def _replace_in_text(match):
        text = match.group(0)
        if 'background-color' in text:
            return text
        return test_name_pattern.sub(lambda m_name: highlight_span, text)

    pre = re.sub(r'>[^<]+<', _replace_in_text, pre)
    return pre + post

print("=== HTML REPLACEMENT TESTS ===")
html_samples = [
    '<h1>TEST NAME</h1>',
    '<div>Test Name: <b>Solidseq Comprehensive Panel On Illumina Novaseq 6000 Platform</b></div>',
    '<div>Test Name: TEST NAME</div>',
    '<p>Indication: Breast and Ovarian Extended Panel - Liquid Biopsy Assay</p>',
    '<p>After End of Report: Whole Exome Sequencing</p>'
]

for hs in html_samples:
    print('INPUT: ', hs)
    print('OUTPUT:', replace_test_name_in_html(hs))
    print()

def set_run_highlight_yellow(r_elem):
    rPr = r_elem.find(qn('w:rPr'))
    if rPr is None:
        rPr = OxmlElement('w:rPr')
        r_elem.insert(0, rPr)
    hl = rPr.find(qn('w:highlight'))
    if hl is None:
        hl = OxmlElement('w:highlight')
        rPr.append(hl)
    hl.set(qn('w:val'), 'yellow')
    
    # Ensure black text for high contrast on yellow
    color = rPr.find(qn('w:color'))
    if color is not None:
        color.set(qn('w:val'), '000000')

def remove_run_highlight(r_elem):
    rPr = r_elem.find(qn('w:rPr'))
    if rPr is not None:
        hl = rPr.find(qn('w:highlight'))
        if hl is not None:
            rPr.remove(hl)

def highlight_test_name_in_docx_obj(doc):
    def highlight_in_paragraph(p_elem, seen_eor=False):
        if seen_eor:
            return
        
        r_elements = [c for c in p_elem if c.tag.endswith('}r')]
        if not r_elements:
            return
        
        full_text = "".join("".join(t.text for t in r.iter() if t.tag.endswith('}t') and t.text) for r in r_elements)
        if not full_text:
            return
        
        # If whole paragraph is strictly TEST NAME or a standalone test title
        cleaned_full = full_text.strip()
        if cleaned_full.upper() == "TEST NAME" or is_test_name_text(cleaned_full):
            for r in r_elements:
                set_run_highlight_yellow(r)
            if cleaned_full.upper() != "TEST NAME":
                t_elems = [t for t in p_elem.iter() if t.tag.endswith('}t')]
                if t_elems:
                    t_elems[0].text = "TEST NAME"
                    t_elems[0].set(qn('xml:space'), 'preserve')
                    for t in t_elems[1:]:
                        t.text = ""
            return

        # Otherwise check individual runs and isolate only the test name
        for r in list(r_elements):
            t_elems = [t for t in r if t.tag.endswith('}t')]
            if not t_elems:
                continue
            r_text = "".join(t.text for t in t_elems if t.text)
            if not r_text:
                continue
            
            match = test_name_pattern.search(r_text)
            if match:
                start, end = match.span()
                matched_str = "TEST NAME"
                before_text = r_text[:start]
                after_text = r_text[end:]
                
                # Clone original rPr BEFORE modifying anything
                rPr_elem = r.find(qn('w:rPr'))
                rPr_copy_before = copy.deepcopy(rPr_elem) if rPr_elem is not None else None
                rPr_copy_after = copy.deepcopy(rPr_elem) if rPr_elem is not None else None
                
                if before_text:
                    t_elems[0].text = before_text
                    for t in t_elems[1:]:
                        t.text = ""
                    remove_run_highlight(r)
                    
                    # New run for highlighted TEST NAME only
                    new_r = OxmlElement('w:r')
                    if rPr_elem is not None:
                        new_r.append(copy.deepcopy(rPr_elem))
                    set_run_highlight_yellow(new_r)
                    new_t = OxmlElement('w:t')
                    new_t.text = matched_str
                    new_t.set(qn('xml:space'), 'preserve')
                    new_r.append(new_t)
                    
                    p_elem.insert(p_elem.index(r) + 1, new_r)
                    
                    if after_text:
                        after_r = OxmlElement('w:r')
                        if rPr_copy_after is not None:
                            after_r.append(rPr_copy_after)
                        remove_run_highlight(after_r)
                        after_t = OxmlElement('w:t')
                        after_t.text = after_text
                        after_t.set(qn('xml:space'), 'preserve')
                        after_r.append(after_t)
                        p_elem.insert(p_elem.index(new_r) + 1, after_r)
                else:
                    t_elems[0].text = matched_str
                    for t in t_elems[1:]:
                        t.text = ""
                    set_run_highlight_yellow(r)
                    
                    if after_text:
                        after_r = OxmlElement('w:r')
                        if rPr_copy_after is not None:
                            after_r.append(rPr_copy_after)
                        remove_run_highlight(after_r)
                        after_t = OxmlElement('w:t')
                        after_t.text = after_text
                        after_t.set(qn('xml:space'), 'preserve')
                        after_r.append(after_t)
                        p_elem.insert(p_elem.index(r) + 1, after_r)

    seen_eor = False
    for p in doc.element.iter():
        if p.tag.endswith('}p'):
            p_text = "".join(t.text for t in p.iter() if t.tag.endswith('}t') and t.text)
            if is_end_of_report_text(p_text):
                seen_eor = True
            highlight_in_paragraph(p, seen_eor=seen_eor)

    for section in doc.sections:
        for h in (section.header, section.first_page_header, section.even_page_header):
            if h is not None:
                for p in h._element.iter():
                    if p.tag.endswith('}p'):
                        highlight_in_paragraph(p, seen_eor=False)
        for f in (section.footer, section.first_page_footer, section.even_page_footer):
            if f is not None:
                for p in f._element.iter():
                    if p.tag.endswith('}p'):
                        highlight_in_paragraph(p, seen_eor=False)

print("=== DOCX HIGHLIGHT TESTS ===")
doc = docx.Document()
p1 = doc.add_paragraph()
p1.add_run("Test Name: Breast and Ovarian Extended Panel - Liquid Biopsy Assay")

p2 = doc.add_paragraph()
p2.add_run("Test Name: ")
p2.add_run("TEST NAME")

p3 = doc.add_paragraph()
p3.add_run("Solidseq Comprehensive Panel On Illumina Novaseq 6000 Platform")

p4 = doc.add_paragraph()
p4.add_run("End of Report")

p5 = doc.add_paragraph()
p5.add_run("Whole Exome Sequencing in methodology after end of report")

highlight_test_name_in_docx_obj(doc)
doc.save("scratch/test_fixed_out.docx")

doc_check = docx.Document("scratch/test_fixed_out.docx")
for idx, p in enumerate(doc_check.paragraphs):
    print(f"P{idx+1}:")
    for r_idx, r in enumerate(p.runs):
        rPr = r._r.find(qn('w:rPr'))
        hl = rPr.find(qn('w:highlight')) if rPr is not None else None
        hl_val = hl.get(qn('w:val')) if hl is not None else None
        print(f"   Run {r_idx+1}: text='{r.text}', highlight={hl_val}")
