import sys
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import docx
from docx.oxml.ns import qn
from converter import convert_json_to_docx, audit_and_heal_docx
import json

sys.stdout.reconfigure(encoding='utf-8')

# Find a json in extracted_jsons
jsons = list(Path("extracted_jsons").glob("*.json"))
if not jsons:
    print("No jsons found in extracted_jsons")
    sys.exit(0)

test_json_p = jsons[0]
print(f"Testing with JSON: {test_json_p.name}")

with open(test_json_p, "r", encoding="utf-8") as f:
    data = json.load(f)

out_docx_path = "output/test_highlight_verification.docx"
convert_json_to_docx(data, output_path=out_docx_path)

doc = docx.Document(out_docx_path)
highlight_count = 0
for i, p in enumerate(doc.paragraphs):
    for r in p.runs:
        rPr = r._r.find(qn('w:rPr'))
        if rPr is not None:
            hl = rPr.find(qn('w:highlight'))
            if hl is not None and hl.get(qn('w:val')) == 'yellow':
                highlight_count += 1
                print(f"Highlighted in p[{i}]: text='{r.text}'")

for ti, t in enumerate(doc.tables):
    for ri, row in enumerate(t.rows):
        for ci, c in enumerate(row.cells):
            for pi, p in enumerate(c.paragraphs):
                for r in p.runs:
                    rPr = r._r.find(qn('w:rPr'))
                    if rPr is not None:
                        hl = rPr.find(qn('w:highlight'))
                        if hl is not None and hl.get(qn('w:val')) == 'yellow':
                            highlight_count += 1
                            print(f"Highlighted in table[{ti}][{ri},{ci}]: text='{r.text}'")

print(f"\n[+] Total yellow highlights found: {highlight_count}")
