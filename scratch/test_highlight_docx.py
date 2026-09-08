import sys
import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path
import re

sys.stdout.reconfigure(encoding='utf-8')

# Let's inspect test names in several docx files
docx_files = list(Path("output").glob("*.docx"))
print(f"Total docx files in output: {len(docx_files)}")

for df in docx_files[:5]:
    try:
        doc = docx.Document(str(df))
        matches = []
        for i, p in enumerate(doc.paragraphs):
            if "TEST NAME" in p.text or "Test Name" in p.text or "Panel" in p.text:
                matches.append((f"p[{i}]", p.text.strip()))
        for ti, t in enumerate(doc.tables):
            for ri, row in enumerate(t.rows):
                for ci, c in enumerate(row.cells):
                    for pi, p in enumerate(c.paragraphs):
                        if "TEST NAME" in p.text or "Test Name" in p.text:
                            matches.append((f"t[{ti}][{ri},{ci}]", p.text.strip()))
        print(f"\n--- {df.name} (matches: {len(matches)}) ---")
        for loc, txt in matches[:5]:
            print(f"  {loc}: {txt[:120]}")
    except Exception as e:
        print(f"Error reading {df.name}: {e}")
