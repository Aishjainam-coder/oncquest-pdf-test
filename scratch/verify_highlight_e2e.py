import json
import os
import sys
sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))
import docx
from docx.oxml.ns import qn
from converter import convert_json_to_docx, generate_dynamic_template_html

def test_conversion():
    json_path = 'extracted_jsons/TestReport_SAROJ DEVI (OQG2604250052)_2600133939_1eb3f888-f97e-4181-b84d-d9536de7af26.json'
    with open(json_path, 'r', encoding='utf-8') as f:
        data = json.load(f)

    docx_bytes = convert_json_to_docx(data)
    out_docx_path = 'scratch/test_saroj_verified.docx'
    with open(out_docx_path, 'wb') as f:
        f.write(docx_bytes)

    doc = docx.Document(out_docx_path)
    docx_hl = []
    for p in doc.paragraphs:
        for r in p.runs:
            rPr = r._r.find(qn('w:rPr'))
            if rPr is not None:
                hl = rPr.find(qn('w:highlight'))
                if hl is not None:
                    docx_hl.append((p.text, r.text, hl.get(qn('w:val'))))

    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for r in p.runs:
                        rPr = r._r.find(qn('w:rPr'))
                        if rPr is not None:
                            hl = rPr.find(qn('w:highlight'))
                            if hl is not None:
                                docx_hl.append((p.text, r.text, hl.get(qn('w:val'))))

    print(f"[+] Total highlighted runs in DOCX: {len(docx_hl)}")
    for p_txt, r_txt, val in docx_hl:
        print(f"    - Highlighted text: '{r_txt}' (color: {val}) in context: '{p_txt[:60]}'")

    # HTML test
    html = generate_dynamic_template_html(data)
    has_html_highlight = ('background-color: #ffff00' in html) or ('background-color: yellow' in html)
    print(f"[+] HTML output contains yellow highlighted test name: {has_html_highlight}")

if __name__ == "__main__":
    test_conversion()
