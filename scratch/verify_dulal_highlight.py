import os
import sys
sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))
import json
import docx
from converter import convert_json_to_docx
from docx.oxml.ns import qn

with open('extracted_jsons/TestReport_DULAL NAHA (KOL2604070057)_2600128556_11d3f985-7e08-40be-91b4-48b35be363ee.json', 'r', encoding='utf-8') as f:
    data = json.load(f)

docx_bytes = convert_json_to_docx(data)
with open('scratch/dulal_test_clean.docx', 'wb') as f:
    f.write(docx_bytes)

doc = docx.Document('scratch/dulal_test_clean.docx')
hl_runs = []
for p in doc.paragraphs:
    for r in p.runs:
        rPr = r._r.find(qn('w:rPr'))
        if rPr is not None and rPr.find(qn('w:highlight')) is not None:
            hl_runs.append((r.text, p.text[:60]))

for t in doc.tables:
    for row in t.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for r in p.runs:
                    rPr = r._r.find(qn('w:rPr'))
                    if rPr is not None and rPr.find(qn('w:highlight')) is not None:
                        hl_runs.append((r.text, p.text[:60]))

print(f"Total highlighted runs in Dulal Naha DOCX: {len(hl_runs)}")
for r_text, ctx in hl_runs:
    print(f"  Highlighted: '{r_text}' in context: '{ctx}'")
