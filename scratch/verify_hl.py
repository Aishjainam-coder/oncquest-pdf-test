import docx
from docx.oxml.ns import qn
import sys

sys.stdout.reconfigure(encoding='utf-8')

doc = docx.Document('scratch/test_birendra_highlighted.docx')
print("--- Paragraphs with highlight ---")
for i, p in enumerate(doc.paragraphs):
    for r in p.runs:
        rPr = r._r.find(qn('w:rPr'))
        if rPr is not None:
            hl = rPr.find(qn('w:highlight'))
            if hl is not None:
                print(f"P[{i}] text={repr(r.text)} highlight={hl.get(qn('w:val'))}")

print("\n--- Tables with highlight ---")
for ti, t in enumerate(doc.tables):
    for ri, row in enumerate(t.rows):
        for ci, c in enumerate(row.cells):
            for pi, p in enumerate(c.paragraphs):
                for r in p.runs:
                    rPr = r._r.find(qn('w:rPr'))
                    if rPr is not None:
                        hl = rPr.find(qn('w:highlight'))
                        if hl is not None:
                            print(f"Table[{ti}][{ri},{ci}] text={repr(r.text)} highlight={hl.get(qn('w:val'))}")
