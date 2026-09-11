import json
import os
import sys
import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import copy
import re

def is_test_name_text(text: str) -> bool:
    if not isinstance(text, str):
        return False
    cleaned = text.strip()
    if len(cleaned) < 4 or len(cleaned) > 130:
        return False
        
    exclude_prefixes = (
        "clinical indication", "sample description", "report highlights", 
        "key findings", "test results", "tier ", "case id", "sample type", 
        "name :", "date & time", "bill. loc", "ref. by", "report version",
        "qr code", "page ", "salient features", "clinical suspicion",
        "dr.", "laboratory", "oncquest", "result summary", "methodology"
    )
    cleaned_lower = cleaned.lower()
    for ex in exclude_prefixes:
        if cleaned_lower.startswith(ex):
            return False
            
    patterns = [
        r'^(?:Breast\s+and\s+Ovarian\s+Extended\s+Panel\s*[-–]\s*Liquid\s+Biopsy\s+Assay)$',
        r'^(?:(?:Liquidseq\s+Actionable|Brainseq)\s+Genomic\s+Profiling\s+Panel(?:\s*[-–]\s*Advance)?)$',
        r'^(?:Whole\s+Exome\s+Sequencing(?:\s+on\s+(?:the\s+)?Illumina\s+[\w\s-]+\s+Platform)?)$',
        r'^[\w\s/&,–\-\(\)\.\+]+?(?:Genomic\s+Profiling\s+Panel|Extended\s+Panel|Profiling\s+Panel|Biopsy\s+Assay|Exome\s+Sequencing|Sequencing\s+Panel|Cancer\s+Panel|Gene\s+Panel|Profiling\s+Assay|Sequencing\s+Assay|Biopsy\s+Panel|NGS\s+Panel)(?:\s*[-–]\s*Advance)?(?:\s*\([^)]*\))?(?:\s+on\s+(?:the\s+)?Illumina\s+[\w\s-]+\s+Platform)?$',
        r'^(?:[A-Z\s]{4,}\s+PANEL(?:\s*[-–]\s*ADVANCE)?(?:\s*\([^)]*\))?)$',
    ]
    for pat in patterns:
        if re.match(pat, cleaned, re.IGNORECASE):
            return True
            
    return False

def is_end_of_report_text(text: str) -> bool:
    if not isinstance(text, str):
        return False
    return bool(re.search(r'end\s+of\s+report', text, re.IGNORECASE))

def set_run_highlight_yellow(r_elem):
    rPr = r_elem.find(qn('w:rPr'))
    if rPr is None:
        rPr = OxmlElement('w:rPr')
        r_elem.insert(0, rPr)
    hl = rPr.find(qn('w:highlight'))
    if hl is None:
        hl = OxmlElement('w:highlight')
        rPr.append(hl)
    hl.set(qn('w:val'), 'yellow')
    
    # Ensure black text for high contrast on yellow
    color = rPr.find(qn('w:color'))
    if color is not None:
        color.set(qn('w:val'), '000000')

test_name_pattern = re.compile(
    r'(\bTEST\s+NAME\b|'
    r'Breast\s+and\s+Ovarian\s+Extended\s+Panel\s*[-–]\s*Liquid\s+Biopsy\s+Assay|'
    r'(?:Liquidseq\s+Actionable|Brainseq)\s+Genomic\s+Profiling\s+Panel(?:\s*[-–]\s*Advance)?|'
    r'Whole\s+Exome\s+Sequencing(?:\s+on\s+(?:the\s+)?Illumina\s+[\w\s-]+\s+Platform)?|'
    r'[\w\s/&,–\-\(\)\.\+]+?(?:Genomic\s+Profiling\s+Panel|Extended\s+Panel|Profiling\s+Panel|Biopsy\s+Assay|Exome\s+Sequencing|Sequencing\s+Panel|Cancer\s+Panel|Gene\s+Panel|Profiling\s+Assay|Sequencing\s+Assay|Biopsy\s+Panel|NGS\s+Panel)(?:\s*[-–]\s*Advance)?(?:\s*\([^)]*\))?(?:\s+on\s+(?:the\s+)?Illumina\s+[\w\s-]+\s+Platform)?)',
    re.IGNORECASE
)

def highlight_test_name_in_docx_obj(doc):
    def highlight_in_paragraph(p_elem, seen_eor=False):
        if seen_eor:
            return
        
        r_elements = [c for c in p_elem if c.tag.endswith('}r')]
        if not r_elements:
            return
        
        full_text = "".join("".join(t.text for t in r.iter() if t.tag.endswith('}t') and t.text) for r in r_elements)
        if not full_text:
            return
        
        # If whole paragraph is TEST NAME or test name pattern
        if full_text.strip().upper() == "TEST NAME" or is_test_name_text(full_text):
            for r in r_elements:
                set_run_highlight_yellow(r)
            if full_text.strip().upper() != "TEST NAME":
                t_elems = [t for t in p_elem.iter() if t.tag.endswith('}t')]
                if t_elems:
                    t_elems[0].text = "TEST NAME"
                    t_elems[0].set(qn('xml:space'), 'preserve')
                    for t in t_elems[1:]:
                        t.text = ""
            return

        # Check individual runs
        for r in list(r_elements):
            t_elems = [t for t in r if t.tag.endswith('}t')]
            if not t_elems:
                continue
            r_text = "".join(t.text for t in t_elems if t.text)
            if not r_text:
                continue
            
            match = test_name_pattern.search(r_text)
            if match:
                start, end = match.span()
                matched_str = "TEST NAME"
                before_text = r_text[:start]
                after_text = r_text[end:]
                
                rPr_orig = r.find(qn('w:rPr'))
                
                if before_text:
                    t_elems[0].text = before_text
                    for t in t_elems[1:]:
                        t.text = ""
                    
                    new_r = OxmlElement('w:r')
                    if rPr_orig is not None:
                        new_r.append(copy.deepcopy(rPr_orig))
                    set_run_highlight_yellow(new_r)
                    new_t = OxmlElement('w:t')
                    new_t.text = matched_str
                    new_t.set(qn('xml:space'), 'preserve')
                    new_r.append(new_t)
                    
                    p_elem.insert(p_elem.index(r) + 1, new_r)
                    
                    if after_text:
                        after_r = OxmlElement('w:r')
                        if rPr_orig is not None:
                            after_r.append(copy.deepcopy(rPr_orig))
                        after_t = OxmlElement('w:t')
                        after_t.text = after_text
                        after_t.set(qn('xml:space'), 'preserve')
                        after_r.append(after_t)
                        p_elem.insert(p_elem.index(new_r) + 1, after_r)
                else:
                    t_elems[0].text = matched_str
                    for t in t_elems[1:]:
                        t.text = ""
                    set_run_highlight_yellow(r)
                    
                    if after_text:
                        after_r = OxmlElement('w:r')
                        if rPr_orig is not None:
                            after_r.append(copy.deepcopy(rPr_orig))
                        after_t = OxmlElement('w:t')
                        after_t.text = after_text
                        after_t.set(qn('xml:space'), 'preserve')
                        after_r.append(after_t)
                        p_elem.insert(p_elem.index(r) + 1, after_r)

    seen_eor = False
    for p in doc.element.iter():
        if p.tag.endswith('}p'):
            p_text = "".join(t.text for t in p.iter() if t.tag.endswith('}t') and t.text)
            if is_end_of_report_text(p_text):
                seen_eor = True
            highlight_in_paragraph(p, seen_eor=seen_eor)

    for section in doc.sections:
        for h in (section.header, section.first_page_header, section.even_page_header):
            if h is not None:
                for p in h._element.iter():
                    if p.tag.endswith('}p'):
                        highlight_in_paragraph(p, seen_eor=False)
        for f in (section.footer, section.first_page_footer, section.even_page_footer):
            if f is not None:
                for p in f._element.iter():
                    if p.tag.endswith('}p'):
                        highlight_in_paragraph(p, seen_eor=False)

# Test creation of docx with standalone and embedded test names
doc = docx.Document()
p1 = doc.add_paragraph()
r1 = p1.add_run("Breast and Ovarian Extended Panel - Liquid Biopsy Assay")

p2 = doc.add_paragraph()
r2 = p2.add_run("This patient was evaluated for Liquidseq Actionable Genomic Profiling Panel with high quality standards.")

p3 = doc.add_paragraph()
r3 = p3.add_run("Normal paragraph with no test name.")

highlight_test_name_in_docx_obj(doc)
doc.save("scratch/test_output_highlight.docx")

# Verify
doc_check = docx.Document("scratch/test_output_highlight.docx")
print("Verification results:")
for i, p in enumerate(doc_check.paragraphs):
    print(f"\nParagraph {i+1}:")
    for j, r in enumerate(p.runs):
        rPr = r._r.find(qn('w:rPr'))
        hl_val = None
        if rPr is not None:
            hl = rPr.find(qn('w:highlight'))
            if hl is not None:
                hl_val = hl.get(qn('w:val'))
        print(f"  Run {j+1}: text='{r.text}', highlight={hl_val}")
