import glob
import os
import docx
from docx.oxml.ns import qn

print("Searching DOCX files:")
for p in glob.glob('output/**/*.docx', recursive=True) + glob.glob('*.docx') + glob.glob('scratch/**/*.docx', recursive=True):
    try:
        doc = docx.Document(p)
        txt = '\n'.join(para.text for para in doc.paragraphs) + '\n'.join(c.text for t in doc.tables for r in t.rows for c in r.cells)
        if '34 unique genes' in txt:
            print(f"\nDOCX found: {p}")
            for para in doc.paragraphs:
                for run in para.runs:
                    rPr = run._r.find(qn('w:rPr'))
                    if rPr is not None and rPr.find(qn('w:highlight')) is not None:
                        hl_val = rPr.find(qn('w:highlight')).get(qn('w:val'))
                        print(f"   PARA RUN HL: '{run.text}' ({hl_val})")
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        for para in cell.paragraphs:
                            for run in para.runs:
                                rPr = run._r.find(qn('w:rPr'))
                                if rPr is not None and rPr.find(qn('w:highlight')) is not None:
                                    hl_val = rPr.find(qn('w:highlight')).get(qn('w:val'))
                                    print(f"   TABLE CELL RUN HL: '{run.text}' ({hl_val})")
    except Exception as e:
        pass

print("\nSearching HTML files:")
for p in glob.glob('output/**/*.html', recursive=True) + glob.glob('*.html') + glob.glob('scratch/**/*.html', recursive=True):
    try:
        with open(p, 'r', encoding='utf-8') as f:
            html = f.read()
        if '34 unique genes' in html:
            print(f"\nHTML found: {p}")
            import re
            for m in re.finditer(r'<[^>]*background(?:-color)?:\s*(?:#ffff00|yellow)[^>]*>(.*?)</[^>]+>', html, re.IGNORECASE):
                print(f"   HTML HL: '{m.group(1)}'")
    except Exception as e:
        pass
